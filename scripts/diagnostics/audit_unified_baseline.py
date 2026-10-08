from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path

from docx import Document


REQUIRED_SECTIONS = [
    "产品定义",
    "项目、输入和来源规则",
    "用户流程和页面",
    "总体系统形态",
    "数据所有权和分层",
    "稳定身份、状态和审计字段",
    "人物身份归一",
    "事件发现、边界和跨章事件",
    "事实版本、状态和角色认知",
    "任务相关检索和上下文选择",
    "分析主张、验证和专项分析",
    "后台任务可靠性",
    "制品、崩溃恢复和局部重算",
    "在线模型、分析方案和提示词",
    "质量、金标和验收标准",
    "参考项目机制总账",
    "当前实现审计",
    "功能与非功能需求基线",
]

REQUIRED_CONTRACTS = [
    "不设置统一小说字数硬上限",
    "只设计、测试和验收 PC/桌面端浏览器",
    "模型提议，程序决定",
    "多匹配进入不确定状态",
    "禁止无保护传递合并",
    "EventCandidate",
    "FactVersion",
    "ActorKnowledge",
    "检索支持证据",
    "检索反面或例外证据",
    "一个任务最多只有一个当前有效尝试",
    "制品 READY + 任务成功 + 尝试成功",
    "读取到 20 个模型",
    "P01～P19 是证据库",
    "FR-017",
    "NFR-011",
]

CURRENT_BASELINE_CONTRACTS = [
    "用户决定方向并拥有最终裁决权",
    "Codex 和后续 AI 只在真实留白处补充",
    "用户拥有最高决策权",
    "Fable 是本项目的总设计师",
    "Codex 和后续接手的 AI 是执行者",
    "概括性指令、口语举例或示例数字",
    "任何新会话、上下文压缩、上下文丢失或新 AI 接手后",
    "不需要价格、币种、金额、单价、费用估算或费用汇总",
    "拆书结果供用户参考学习，不直接交给写作 Agent",
    "拆书结果不自动生成用户新书的最小开书包",
    "42 问全量审计与共享能力包",
    "七问组合不是 Fable 固定批次",
    "逐问增量答案编译已实现",
    "首组五问用户反馈与基础修复闭环",
    "2.1 合同闭环、长输入最小诊断",
    "4.9 只负责章末钩类型配比、轮换和强弱节律",
    "学习答案每问一个独立模型任务",
    "学习页默认一次只展示一问",
    "独立 3.4 兑现追踪",
    "2.1 合同完整性纠偏与长输入最小诊断",
    "合同—原料—返回结构—程序验收",
    "精简相关输入 vs 约 16 万 Token 长输入",
    "15 万 Token 默认软上限",
    "2.1 精简臂真实复测与本机长任务传输纠偏",
    "输入 79,582 Token、输出 15,892 Token",
    "输入 165,478 Token、输出 19,323 Token",
    "输入 164,030 Token、输出 3,078 Token",
    "结果 24/24",
    "自由正文复核纠偏",
    "模型提议、程序决定落地",
    "创作学习报告）即使走本机也强制流式",
    "统一软上限、2.1 最小模型合同与独立 3.4",
    "所有已配置模型和核心 42 问",
    "输入 73,297 Token、输出 3,485 Token",
    "总输入 210,476 Token",
    "其余 8 个待运行任务自动取消",
    "外部数据边界程序化",
    "刷新首组中因合同/专项版本变化而过期的 1.4、2.2、4.9",
]

README_CURRENT_CONTRACTS = [
    "任何新会话、上下文压缩或后续 AI 接手后",
    "首组反馈基础修复闭环已经完成",
    "数百至上千章按连续窗口覆盖",
    "每问一个独立模型任务",
    "学习页仍按八阶段目录一次只读一问",
    "2.1 已成为第一个具备逐项合同验收的问题",
    "输入 73,297 Token",
    "113 人不重不漏",
    "所有已配置模型和核心 42 问",
    "总输入 210,476 Token",
    "其余 8 个自动取消",
    "24/24",
    "均由程序编译",
    "本机学习报告等长任务也强制流式",
    "15 万 Token 单次请求软上限",
    "下一步是用新合同刷新已过期的 1.4、2.2、4.9",
]

README_FORBIDDEN_CURRENT_CLAIMS = [
    "共 `5/7` 项专项原料就绪",
    "唯一下一步是补齐 5.3",
    "正式学习答案仍为 `0/42`",
    "16k/24k/32k/48k/64k",
    "本轮实现没有发起新的正式模型调用",
    "当前只剩同问同证据的约 16 万 Token 长输入臂",
    "约 15 万 Token 只是当前模型的候选安全上限",
]


def docx_text(path: Path) -> str:
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(blocks)


def check_required(text: str, *, label: str) -> None:
    missing_sections = [value for value in REQUIRED_SECTIONS if value not in text]
    missing_contracts = [value for value in REQUIRED_CONTRACTS if value not in text]
    if missing_sections or missing_contracts:
        raise SystemExit(
            f"{label} 缺少内容：sections={missing_sections}, contracts={missing_contracts}"
        )


def check_values(text: str, values: list[str], *, label: str) -> None:
    missing = [value for value in values if value not in text]
    if missing:
        raise SystemExit(f"{label} 缺少当前决策：{missing}")


def check_absent(text: str, values: list[str], *, label: str) -> None:
    present = [value for value in values if value in text]
    if present:
        raise SystemExit(f"{label} 仍包含已撤回的当前决策：{present}")


def check_core_question_coverage(markdown: str) -> None:
    core_section = re.search(
        r"共 66 问.*?(?P<table>\| 编号 \| 核心问题 \|.*?)(?=\n#### 2\.2\.2)",
        markdown,
        flags=re.DOTALL,
    )
    package_section = re.search(
        r"#### 2\.2\.2 42 问全量审计与共享能力包.*?"
        r"(?P<table>\| 能力包 \| 主问题.*?)(?=\n#### 2\.2\.3)",
        markdown,
        flags=re.DOTALL,
    )
    if core_section is None or package_section is None:
        raise SystemExit("Markdown 无法定位核心问题表或共享能力包表。")

    core_questions = re.findall(
        r"^\| (\d+\.\d+) \|", core_section.group("table"), flags=re.MULTILINE
    )
    package_rows = re.findall(
        r"^\| (QG-\d+) [^|]*\| ([^|]+)\|",
        package_section.group("table"),
        flags=re.MULTILINE,
    )
    package_labels = [label for label, _ in package_rows]
    packaged_questions = [
        question
        for _, values in package_rows
        for question in re.findall(r"`(\d+\.\d+)`", values)
    ]

    core_counts = Counter(core_questions)
    package_counts = Counter(packaged_questions)
    expected_labels = [f"QG-{index}" for index in range(1, 9)]
    problems: list[str] = []
    if len(core_questions) != 42 or any(count != 1 for count in core_counts.values()):
        problems.append(
            f"核心问题表应有 42 个唯一编号，实际共 {len(core_questions)} 项、"
            f"唯一 {len(core_counts)} 项"
        )
    if package_labels != expected_labels:
        problems.append(f"能力包应依次为 {expected_labels}，实际为 {package_labels}")
    missing = sorted(core_counts.keys() - package_counts.keys())
    extra = sorted(package_counts.keys() - core_counts.keys())
    repeated = sorted(
        question for question, count in package_counts.items() if count != 1
    )
    if missing or extra or repeated or len(packaged_questions) != 42:
        problems.append(
            "能力包必须不重不漏覆盖核心 42 问："
            f"总数={len(packaged_questions)}, missing={missing}, "
            f"extra={extra}, repeated={repeated}"
        )
    if problems:
        raise SystemExit("；".join(problems))


def main() -> None:
    parser = argparse.ArgumentParser(description="检查统一基线关键内容和旧文档清理状态。")
    parser.add_argument("--repo", type=Path, required=True)
    parser.add_argument("--research", type=Path, required=True)
    parser.add_argument("--check-cleanup", action="store_true")
    args = parser.parse_args()

    source = args.repo / "docs" / "CURRENT_BASELINE.md"
    published = args.research / "00_当前设计" / "AI小说拆解工作台_统一产品与系统设计基线_V1.0.docx"

    markdown = source.read_text(encoding="utf-8")
    check_required(markdown, label="Markdown")
    check_values(markdown, CURRENT_BASELINE_CONTRACTS, label="Markdown")
    check_core_question_coverage(markdown)

    if published.exists():
        published_text = docx_text(published)
        check_required(published_text, label="Word")
        check_values(published_text, CURRENT_BASELINE_CONTRACTS, label="Word")

    readme = (args.repo / "README.md").read_text(encoding="utf-8")
    check_values(readme, README_CURRENT_CONTRACTS, label="README")
    check_absent(readme, README_FORBIDDEN_CURRENT_CLAIMS, label="README")

    if args.check_cleanup:
        forbidden_archives = [
            args.research / "90_历史归档",
            args.repo / "docs" / "archive",
        ]
        leftovers = [str(path) for path in forbidden_archives if path.exists()]
        if leftovers:
            raise SystemExit(f"不应建立历史文档归档目录：{leftovers}")

    print("统一基线内容与清理状态检查通过。")


if __name__ == "__main__":
    main()
