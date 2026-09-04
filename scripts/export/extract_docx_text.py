from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def extract(path: Path) -> str:
    document = Document(path)
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading"):
            try:
                level = int(style.rsplit(" ", 1)[1])
            except (IndexError, ValueError):
                level = 1
            blocks.append(f"{'#' * max(1, min(level, 6))} {text}")
        else:
            blocks.append(text)

    for index, table in enumerate(document.tables, start=1):
        blocks.append(f"## [表格 {index}]")
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))

    return "\n\n".join(blocks) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description="结构化提取 DOCX 正文、标题和表格。")
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    result = extract(args.input)
    if args.output is None:
        print(result, end="")
        return

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(result, encoding="utf-8")
    print(args.output)


if __name__ == "__main__":
    main()
