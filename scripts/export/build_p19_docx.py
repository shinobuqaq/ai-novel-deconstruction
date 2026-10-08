from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
REFERENCE_DIR = ROOT / "docs" / "research" / "reference_projects"
SOURCE = REFERENCE_DIR / "AI小说拆书项目评估档案19_SillyTavern核心机制与可用零件审计_V2.0.md"
OUTPUT = REFERENCE_DIR / "AI小说拆书项目评估档案19_SillyTavern核心机制与可用零件审计_V2.0.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
HEADER_FILL = "E8EEF5"
GRID = "B8C2CC"
MUTED = RGBColor(95, 105, 115)


def set_font(run, name: str = "Calibri", size: float | None = None, *, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        row.height = None
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), GRID)


def add_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element

    def create(abstract_id: int, num_id: int, fmt: str, text: str, font: str | None = None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        level.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        level.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        level.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.append(ind)
        level.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            r_pr.append(fonts)
            level.append(r_pr)
        abstract.append(level)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)

    create(70, 70, "bullet", "•", "Symbol")
    create(71, 71, "decimal", "%1.")
    return 70, 71


def apply_num(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)


def add_rich_text(paragraph, text: str):
    parts = re.split(r"(`[^`]+`|https?://\S+)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Consolas", 9.5, color=RGBColor(45, 55, 65))
        else:
            run = paragraph.add_run(part)
            set_font(run, "Calibri", 11)


def add_body_paragraph(doc: Document, text: str, num_id: int | None = None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6 if num_id is None else 4)
    paragraph.paragraph_format.line_spacing = 1.25
    if num_id is not None:
        apply_num(paragraph, num_id)
    add_rich_text(paragraph, text)
    return paragraph


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    return [rows[0], *rows[2:]]


def add_table(doc: Document, rows: list[list[str]]):
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    if column_count == 2:
        widths = [2700, 6660]
    elif column_count == 3:
        widths = [2100, 3300, 3960]
    elif column_count == 4:
        widths = [2050, 2300, 1650, 3360]
    else:
        widths = [9360 // column_count] * column_count
        widths[-1] += 9360 - sum(widths)
    for row_index, values in enumerate(rows):
        for col_index in range(column_count):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            value = values[col_index] if col_index < len(values) else ""
            add_rich_text(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if row_index == 0:
                    run.bold = True
            if row_index == 0:
                shade(cell, HEADER_FILL)
    set_table_geometry(table, widths)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    run = hp.add_run("AI 小说拆书项目评估档案 19")
    set_font(run, size=9, bold=True, color=MUTED)
    run = hp.add_run("  |  P19 · V2.0")
    set_font(run, size=9, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    run = fp.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    run = fp.add_run(" 页")
    set_font(run, size=9, color=MUTED)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    bullet_id, decimal_id = add_numbering(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("AI 小说拆书项目评估档案 19")
    set_font(run, size=24, bold=True, color=RGBColor(25, 45, 70))

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("SillyTavern 核心机制与可用零件审计")
    set_font(run, size=14, bold=True, color=RGBColor.from_string(BLUE))

    metadata = doc.add_table(rows=5, cols=2)
    meta_rows = [
        ("版本", "V2.0"),
        ("研究日期", "2026-07-18"),
        ("研究对象", "SillyTavern 1.18.0"),
        ("固定源码版本", "8172dcd0ee672d3cd9a5e5f7af134f91a45cd2b8"),
        ("许可证", "AGPL-3.0"),
    ]
    for index, (label, value) in enumerate(meta_rows):
        metadata.cell(index, 0).text = label
        metadata.cell(index, 1).text = value
        for run in metadata.cell(index, 0).paragraphs[0].runs:
            set_font(run, size=9.5, bold=True)
        for run in metadata.cell(index, 1).paragraphs[0].runs:
            set_font(run, size=9.5)
        shade(metadata.cell(index, 0), HEADER_FILL)
    set_table_geometry(metadata, [2700, 6660])
    set_table_borders(metadata)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Skip the Markdown title block already represented above.
    index = 0
    while index < len(lines) and not lines[index].startswith("## 1."):
        index += 1

    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and lines[index + 1].startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, parse_table(table_lines))
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
        elif re.match(r"^\d+\. ", line):
            add_body_paragraph(doc, re.sub(r"^\d+\. ", "", line), decimal_id)
        elif line.startswith("- "):
            add_body_paragraph(doc, line[2:], bullet_id)
        else:
            add_body_paragraph(doc, line)
        index += 1

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
