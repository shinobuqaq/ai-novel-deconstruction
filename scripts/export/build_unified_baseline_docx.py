from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from build_p19_docx import (
    BLUE,
    HEADER_FILL,
    add_body_paragraph,
    add_numbering,
    add_table,
    configure_document,
    parse_table,
    set_font,
    set_table_borders,
    set_table_geometry,
    shade,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SOURCE = PROJECT_ROOT / "docs" / "CURRENT_BASELINE.md"
OUTPUT_DIR = PROJECT_ROOT / "docs" / "export"
OUTPUT = OUTPUT_DIR / "AI小说拆解工作台_统一产品与系统设计基线_V1.0.docx"


def build():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    updated_at = next(
        (
            line.split("：", 1)[1].strip()
            for line in lines
            if line.startswith("- 更新时间：")
        ),
        "未注明",
    )
    doc = Document()
    configure_document(doc)
    bullet_id, decimal_id = add_numbering(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("AI 小说拆解工作台")
    set_font(run, size=26, bold=True, color=RGBColor(25, 45, 70))

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("统一产品与系统设计基线 V1.0")
    set_font(run, size=15, bold=True, color=RGBColor.from_string(BLUE))

    metadata = doc.add_table(rows=5, cols=2)
    meta_rows = [
        ("版本", "V1.0"),
        ("更新时间", updated_at),
        ("文档性质", "当前唯一产品与系统判断入口"),
        ("主要读者", "产品设计者、开发者、后续接手的 AI"),
        ("内容源", "代码仓库 docs/CURRENT_BASELINE.md"),
    ]
    for index, (label, value) in enumerate(meta_rows):
        metadata.cell(index, 0).text = label
        metadata.cell(index, 1).text = value
        for run in metadata.cell(index, 0).paragraphs[0].runs:
            set_font(run, size=9.5, bold=True)
        for run in metadata.cell(index, 1).paragraphs[0].runs:
            set_font(run, size=9.5)
        shade(metadata.cell(index, 0), HEADER_FILL)
    set_table_geometry(metadata, [2700, 6660])
    set_table_borders(metadata)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    index = 0
    while index < len(lines) and not lines[index].startswith(
        "## 0. 接手项目时必须先执行的决策程序"
    ):
        index += 1
    if index >= len(lines):
        raise RuntimeError("当前基线缺少第 0 章决策程序，拒绝生成不完整 Word 版")

    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and lines[index + 1].startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, parse_table(table_lines))
            continue
        if line.startswith("```text"):
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(12)
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run("\n".join(code_lines))
            set_font(run, "Consolas", 9.5, color=RGBColor(45, 55, 65))
        elif line.startswith("#### "):
            doc.add_paragraph(line[5:], style="Heading 3")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
        elif re.match(r"^\d+\. ", line):
            add_body_paragraph(doc, re.sub(r"^\d+\. ", "", line), decimal_id)
        elif line.startswith("- "):
            add_body_paragraph(doc, line[2:], bullet_id)
        else:
            add_body_paragraph(doc, line)
        index += 1

    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("AI 小说拆解工作台  |  统一产品与系统设计基线 V1.0")
    set_font(run, size=9, color=RGBColor(95, 105, 115))

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
